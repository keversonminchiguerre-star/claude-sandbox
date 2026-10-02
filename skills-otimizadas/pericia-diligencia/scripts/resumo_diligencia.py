# -*- coding: utf-8 -*-
"""Gera o Resumo para Diligência (docx).
Uso: python3 resumo_diligencia.py proc.json "Resumo para Diligência - <processo>.docx"
ou: from resumo_diligencia import build_resumo; build_resumo(saida, proc)"""
import json
import sys

from docx import Document
from docx.shared import Pt, Emu, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Verdana"


def set_margins(doc, top=1.8, bottom=1.8, left=2.0, right=1.8):
    for section in doc.sections:
        section.page_width = Emu(7560310)    # A4
        section.page_height = Emu(10692130)
        section.top_margin = Cm(top); section.bottom_margin = Cm(bottom)
        section.left_margin = Cm(left); section.right_margin = Cm(right)


def add_p(doc, text="", bold=False, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
          space_after=4, space_before=0, italic=False):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(space_before); pf.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = FONT; r.font.size = Pt(size); r.bold = bold; r.italic = italic
    return p


def _bottom_border(p, sz, color):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), sz)
    b.set(qn('w:space'), '1'); b.set(qn('w:color'), color)
    pBdr.append(b); pPr.append(pBdr)


def add_rule(doc, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    _bottom_border(p, '6', '808080')
    return p


def add_section_header(doc, text, size=11.5, space_before=12, space_after=6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(space_before); pf.space_after = Pt(space_after); pf.line_spacing = 1.0
    r = p.add_run(text.upper())
    r.font.name = FONT; r.font.size = Pt(size); r.bold = True
    _bottom_border(p, '4', 'A6A6A6')
    return p


def add_field(doc, label, value, size=11, space_after=9):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(space_after); pf.space_before = Pt(0); pf.line_spacing = 1.15
    r1 = p.add_run(label + " ")
    r1.font.name = FONT; r1.font.size = Pt(size); r1.bold = True
    r2 = p.add_run(value)
    r2.font.name = FONT; r2.font.size = Pt(size); r2.bold = False
    return p


def build_resumo(out_path, proc):
    doc = Document(); set_margins(doc)
    add_p(doc, "RESUMO PARA DILIGÊNCIA PERICIAL", bold=True, size=13.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p(doc, f"Processo nº {proc['numero']}", size=10.5,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_rule(doc, space_after=10)

    add_section_header(doc, "Identificação do processo")
    for lbl, val in proc["partes"]:
        add_field(doc, lbl, val)
    add_field(doc, "Objetivo da perícia:", proc["objetivo"])

    add_section_header(doc, "Diligência")
    add_field(doc, "Data, horário e local:", proc["diligencia"])

    add_section_header(doc, "Histórico contratual")
    add_field(doc, "Admissão/demissão e evolução de cargos e funções:", proc["admissao"])

    add_section_header(doc, "Síntese das alegações das partes")
    add_field(doc, "Reclamante:", proc["sintese_inicial"])
    add_field(doc, "Reclamada(s):", proc["sintese_contestacao"])

    add_section_header(doc, "Agentes, documentação e EPI")
    add_field(doc, "Agente(s)/atividade(s) questionados:", proc["agentes"])
    add_field(doc, "Documentação apresentada:", proc["documentos"])
    if proc.get("documentos_ausentes"):
        add_field(doc, "Documentação não localizada até o momento:", proc["documentos_ausentes"])
    add_field(doc, "EPI:", proc["epi"])

    add_section_header(doc, "Pontos de atenção para a diligência")
    add_field(doc, "Particularidade(s) a destacar:", proc["particularidade"])
    add_field(doc, "Ponto(s) controvertido(s) quanto à matéria da perícia:",
              proc["controvertido"], space_after=0)

    doc.save(out_path)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build_resumo(sys.argv[2], json.load(f))
